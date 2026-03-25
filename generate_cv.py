"""
CV Design Nude/Beige — Conseillère en Insertion Professionnelle
Louiza Hadid — Mission Locale du Pays de Lorient
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Palette ───────────────────────────────────────────────────────────────────
HDR_BG  = "C4A882"   # header — beige doré
SIDE_BG = "EDE5D8"   # sidebar — beige clair
BODY_BG = "FDFAF6"   # colonne droite — crème
BROWN   = RGBColor(110, 82, 50)     # titres sections
WHT     = RGBColor(255, 255, 255)   # texte sur fond foncé
CREAM   = RGBColor(255, 244, 228)   # sous-titre header
DARK    = RGBColor(45, 42, 38)      # corps de texte
GRAY    = RGBColor(118, 110, 98)    # texte secondaire

FONT   = "Calibri"
OUTPUT = "/home/user/Lettre-de-motivation/CV_Louiza_Hadid_CIP.docx"


# ── Helpers génériques ────────────────────────────────────────────────────────

def mk_run(p, text, size=9.5, bold=False, italic=False, color=None):
    r = p.add_run(text)
    r.font.name      = FONT
    r.font.size      = Pt(size)
    r.font.bold      = bold
    r.font.italic    = italic
    r.font.color.rgb = color or DARK
    return r


def _clr(el, tag):
    bd = OxmlElement(tag)
    for s in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{s}")
        b.set(qn("w:val"), "none"); b.set(qn("w:sz"), "0")
        b.set(qn("w:space"), "0"); b.set(qn("w:color"), "auto")
        bd.append(b)
    ex = el.find(qn(tag))
    if ex is not None:
        el.remove(ex)
    el.append(bd)


def bl_tbl(doc, rows, cols):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = "Normal Table"
    tP = t._tbl.find(qn("w:tblPr"))
    if tP is None:
        tP = OxmlElement("w:tblPr"); t._tbl.insert(0, tP)
    _clr(tP, "w:tblBorders")
    for row in t.rows:
        for cell in row.cells:
            cP = cell._tc.get_or_add_tcPr()
            _clr(cP, "w:tcBorders")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for p in list(cell.paragraphs):
                p._element.getparent().remove(p._element)
    return t


def col_w(cell, cm):
    cP = cell._tc.get_or_add_tcPr()
    w = OxmlElement("w:tcW")
    w.set(qn("w:w"), str(int(cm * 567)))
    w.set(qn("w:type"), "dxa")
    ex = cP.find(qn("w:tcW"))
    if ex is not None:
        cP.remove(ex)
    cP.append(w)


def cell_bg(cell, hx):
    cP = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hx)
    ex = cP.find(qn("w:shd"))
    if ex is not None:
        cP.remove(ex)
    cP.append(shd)


def cell_pad(cell, top=0.15, bot=0.15, left=0.35, right=0.3):
    cP = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for name, val in [("top", top), ("bottom", bot),
                      ("left", left), ("right", right)]:
        m = OxmlElement(f"w:{name}")
        m.set(qn("w:w"), str(int(val * 567)))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    ex = cP.find(qn("w:tcMar"))
    if ex is not None:
        cP.remove(ex)
    cP.append(mar)


def sec_title(container, text, side=False):
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after  = Pt(4)
    r = p.add_run(text.upper())
    r.font.name = FONT
    r.font.size = Pt(8.5 if side else 9.5)
    r.font.bold = True
    r.font.color.rgb = BROWN
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "4")
    bot.set(qn("w:space"), "1");    bot.set(qn("w:color"), "C4A882")
    pBdr.append(bot); pPr.append(pBdr)


def side_item(cell, text, size=9, sb=0, sa=2, indent=False, color=None,
              bold=False, italic=False):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after  = Pt(sa)
    if indent:
        p.paragraph_format.left_indent = Cm(0.15)
    mk_run(p, text, size=size, color=color, bold=bold, italic=italic)
    return p


def exp_block(cm_cell, poste, structure, lieu, dates, bullets):
    p = cm_cell.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(0)
    mk_run(p, poste, size=9.5, bold=True)
    mk_run(p, f"  \u2014  {dates}", size=9, color=GRAY)

    p2 = cm_cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after  = Pt(2)
    mk_run(p2, f"{structure} \u2014 {lieu}", size=9, italic=True, color=GRAY)

    for b in bullets:
        pb = cm_cell.add_paragraph()
        pb.paragraph_format.space_before = Pt(0)
        pb.paragraph_format.space_after  = Pt(1.5)
        pb.paragraph_format.left_indent  = Cm(0.3)
        mk_run(pb, f"\u2013 {b}", size=9.5)


# ═══════════════════════════════════════════════════════════════════════════════
#  Document
# ═══════════════════════════════════════════════════════════════════════════════
doc = Document()
sec = doc.sections[0]
sec.top_margin    = Cm(1.5)
sec.bottom_margin = Cm(1.5)
sec.left_margin   = Cm(1.5)
sec.right_margin  = Cm(1.5)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(9.5)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after  = Pt(0)


# ── 1. HEADER ─────────────────────────────────────────────────────────────────
hdr = bl_tbl(doc, 1, 1)
hc = hdr.cell(0, 0)
col_w(hc, 18.0)
cell_bg(hc, HDR_BG)
cell_pad(hc, top=0.45, bot=0.4, left=0.55, right=0.55)

p = hc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after  = Pt(4)
mk_run(p, "Louiza Hadid", size=24, bold=True, color=WHT)

p2 = hc.add_paragraph()
p2.paragraph_format.space_before = Pt(0)
p2.paragraph_format.space_after  = Pt(10)
mk_run(p2, "Conseillère en insertion professionnelle", size=12, italic=True, color=CREAM)

p3 = hc.add_paragraph()
p3.paragraph_format.space_before = Pt(0)
p3.paragraph_format.space_after  = Pt(0)
mk_run(p3,
    "06 50 37 56 47  \u2022  louizahdd@gmail.com  \u2022  "
    "3 rue de Ventspils, 56100 Lorient  \u2022  Permis B",
    size=9, color=RGBColor(240, 228, 210))


# ── 2. BODY — 2 colonnes ──────────────────────────────────────────────────────
body = bl_tbl(doc, 1, 2)
cs = body.cell(0, 0)   # sidebar
cm = body.cell(0, 1)   # main

col_w(cs, 6.2)
col_w(cm, 11.8)
cell_bg(cs, SIDE_BG)
cell_bg(cm, BODY_BG)
cell_pad(cs, top=0.25, bot=0.25, left=0.4, right=0.3)
cell_pad(cm, top=0.25, bot=0.25, left=0.45, right=0.4)


# ════════════════════════════════════════════════
#  SIDEBAR (gauche)
# ════════════════════════════════════════════════

# — Formation —
sec_title(cs, "Formation", side=True)

side_item(cs, "Licence Sciences de l'\u00c9ducation", size=9, bold=False, sb=3, sa=1)
# hack: bold separately
p = cs.paragraphs[-1]
for r in p.runs:
    r.font.bold = True

side_item(cs, "Universit\u00e9 Rennes 2", size=8.5, sa=0, color=GRAY)

p = cs.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after  = Pt(4)
mk_run(p, "2022 \u2013 2025", size=8.5, italic=True, color=GRAY)

side_item(cs, "Option : Action \u00c9ducation", size=8.5, sb=0, sa=1, indent=True)
side_item(cs, "et Formation (2 semestres)", size=8.5, sb=0, sa=3, indent=True)

for ligne in [
    "\u00b7 Dispositifs d'insertion (CEJ, CIVIS, apprentissage)",
    "\u00b7 Accompagnement des publics",
    "\u00b7 Psychologie du d\u00e9veloppement",
    "\u00b7 Enqu\u00eates terrain : observations classe,",
    "  entretiens enseignants & \u00e9l\u00e8ves",
]:
    side_item(cs, ligne, size=8.5, sa=1, indent=True)

p_sp = cs.add_paragraph()
p_sp.paragraph_format.space_before = Pt(5)
p_sp.paragraph_format.space_after  = Pt(1)
mk_run(p_sp, "Baccalaur\u00e9at G\u00e9n\u00e9ral", size=9, bold=True)

side_item(cs, "Lyc\u00e9e Dupuy de L\u00f4me \u2014 Lorient", size=8.5, sa=0, color=GRAY)
side_item(cs, "2018 \u2013 2021  \u00b7  SES \u00b7 LLCER Anglais", size=8.5, sa=0,
          italic=True, color=GRAY)

# — Compétences —
sec_title(cs, "Comp\u00e9tences", side=True)
for sk in [
    "\u00b7 \u00c9coute active et bienveillance",
    "\u00b7 Accompagnement individualis\u00e9",
    "\u00b7 Techniques d\u2019entretien",
    "\u00b7 Suivi administratif de dossiers",
    "\u00b7 Animation de groupe",
    "\u00b7 M\u00e9diation et relationnel",
    "\u00b7 Travail en \u00e9quipe",
    "\u00b7 Rigueur et organisation",
]:
    side_item(cs, sk, size=9, sa=2)

# — Langues & Outils —
sec_title(cs, "Langues & Outils", side=True)
for item in ["\u00b7 Fran\u00e7ais \u2014 natif", "\u00b7 Anglais \u2014 C1"]:
    side_item(cs, item, size=9, sa=2)

p_inf = cs.add_paragraph()
p_inf.paragraph_format.space_before = Pt(5)
p_inf.paragraph_format.space_after  = Pt(2)
mk_run(p_inf, "Informatique", size=9, bold=True)
for item in ["\u00b7 Pack Office (Word, Excel)", "\u00b7 Outils collaboratifs en ligne"]:
    side_item(cs, item, size=9, sa=2)


# ════════════════════════════════════════════════
#  MAIN COLUMN (droite)
# ════════════════════════════════════════════════

# — Profil —
sec_title(cm, "Profil")
p = cm.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after  = Pt(0)
mk_run(p,
    "Dipl\u00f4m\u00e9e d\u2019une Licence Sciences de l\u2019\u00c9ducation, je suis rigoureuse, "
    "organis\u00e9e et dot\u00e9e d\u2019un sens aigu de l\u2019\u00e9coute. "
    "Mes exp\u00e9riences dans l\u2019animation en milieu prioritaire, la relation client "
    "et le tutorat de publics allophones m\u2019ont appris \u00e0 adapter mon approche, "
    "\u00e0 g\u00e9rer des situations complexes et \u00e0 assurer un suivi individualis\u00e9. "
    "Familiarise\u0301e avec les dispositifs d\u2019insertion gr\u00e2ce \u00e0 ma formation, "
    "je suis motiv\u00e9e \u00e0 mettre mes comp\u00e9tences au service des jeunes "
    "de la Mission Locale du Pays de Lorient.",
    size=9.5)

# — Expériences —
sec_title(cm, "Exp\u00e9riences professionnelles")

exp_block(cm,
    "T\u00e9l\u00e9op\u00e9ratrice Recouvrement & Auditrice",
    "T\u00e9l\u00e9maque \u2013 Groupe P\u00e9n\u00e9lope", "Caudan (56)",
    "Sept. 2025 \u2013 pr\u00e9sent",
    [
        "Conduite d\u2019appels entrants/sortants : \u00e9coute active, reformulation, recherche de solutions",
        "Appels myst\u00e8res : \u00e9valuation qualit\u00e9 d\u2019accueil, grille de notation, reporting",
        "Suivi administratif rigoureux des dossiers (saisie, mise \u00e0 jour, classement)",
    ]
)

exp_block(cm,
    "Charg\u00e9e de Relation Client",
    "Foncia", "Rennes",
    "Mai \u2013 sept. 2025",
    [
        "Accueil physique et t\u00e9l\u00e9phonique, orientation et traitement des demandes",
        "Suivi individualis\u00e9 des dossiers, r\u00e9daction de comptes rendus, Pack Office quotidien",
        "Gestion de situations parfois conflictuelles avec calme et professionnalisme",
    ]
)

exp_block(cm,
    "B\u00e9n\u00e9vole \u2014 Tutorat de fran\u00e7ais pour allophones",
    "Association Sup\u00c9ducation", "Rennes",
    "Janv. \u2013 avr. 2024",
    [
        "Accompagnement individualis\u00e9 en petits groupes (4\u20136 pers.) \u00e0 distance",
        "Adaptation p\u00e9dagogique aux besoins de chaque apprenant, suivi de la progression",
        "D\u00e9veloppement de l\u2019\u00e9coute active, de la patience et de la m\u00e9diation",
    ]
)

exp_block(cm,
    "Animatrice P\u00e9riscolaire",
    "Mairie de Rennes \u2013 \u00c9cole REP", "Rennes",
    "Janv. \u2013 mai 2023",
    [
        "Accueil et encadrement d\u2019enfants en r\u00e9seau d\u2019\u00e9ducation prioritaire",
        "Adaptation aux publics diversifi\u00e9s, maintien d\u2019un cadre bienveillant et structurant",
        "Collaboration avec l\u2019\u00e9quipe p\u00e9dagogique et lien avec les familles",
    ]
)

exp_block(cm,
    "\u00c9quipi\u00e8re \u2192 Formatrice \u2192 Chef d\u2019\u00e9quipe",
    "McDonald\u2019s", "Rennes",
    "Sept. 2021 \u2013 nov. 2022",
    [
        "Formation et int\u00e9gration des nouveaux collaborateurs \u00e0 leurs postes",
        "Coordination et management d\u2019\u00e9quipe, gestion des priorit\u00e9s sous pression",
        "Progression rapide : \u00e9quipi\u00e8re \u2192 formatrice \u2192 chef d\u2019\u00e9quipe en 14 mois",
    ]
)


# ── Sauvegarde ────────────────────────────────────────────────────────────────
doc.save(OUTPUT)
print(f"Fichier g\u00e9n\u00e9r\u00e9 : {OUTPUT}")
