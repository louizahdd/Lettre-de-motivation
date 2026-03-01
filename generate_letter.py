from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Calibri"

doc = Document()

# ── Marges ──────────────────────────────────────────────────────────────────
section = doc.sections[0]
section.top_margin    = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin   = Cm(2.5)
section.right_margin  = Cm(2.5)

# Supprimer l'espacement par défaut du style Normal
style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = Pt(11)
style.paragraph_format.space_before = Pt(0)
style.paragraph_format.space_after  = Pt(0)
style.paragraph_format.line_spacing = Pt(14)


# ── Helpers ──────────────────────────────────────────────────────────────────
def font(run, size=11, bold=False, italic=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def para(doc, text="", align=WD_ALIGN_PARAGRAPH.LEFT,
         bold=False, size=11, sb=0, sa=6, indent=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after  = Pt(sa)
    if indent:
        p.paragraph_format.first_line_indent = Cm(1)
    if text:
        r = p.add_run(text)
        font(r, size=size, bold=bold)
    return p


def remove_borders(obj_xml, sides=("top","left","bottom","right","insideH","insideV"), tag="w:tcBorders"):
    borders_el = OxmlElement(tag)
    for side in sides:
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"),   "none")
        b.set(qn("w:sz"),    "0")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "auto")
        borders_el.append(b)
    existing = obj_xml.find(qn(tag))
    if existing is not None:
        obj_xml.remove(existing)
    obj_xml.append(borders_el)


def set_cell_width(cell, width_cm):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW  = OxmlElement("w:tcW")
    tcW.set(qn("w:w"),    str(int(width_cm * 567)))  # twips : 1 cm ≈ 567
    tcW.set(qn("w:type"), "dxa")
    existing = tcPr.find(qn("w:tcW"))
    if existing is not None:
        tcPr.remove(existing)
    tcPr.append(tcW)


def add_cell_para(cell, text, bold=False, size=11,
                  align=WD_ALIGN_PARAGRAPH.LEFT, sa=3):
    p = cell.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(sa)
    r = p.add_run(text)
    font(r, size=size, bold=bold)
    return p


# ════════════════════════════════════════════════════════════════════════════
#  EN-TÊTE
# ════════════════════════════════════════════════════════════════════════════
table = doc.add_table(rows=1, cols=2)
table.style = "Normal Table"   # pas de bordures par défaut

# Supprimer aussi les bordures au niveau du tableau
tbl    = table._tbl
tblPr  = tbl.find(qn("w:tblPr"))
if tblPr is None:
    tblPr = OxmlElement("w:tblPr")
    tbl.insert(0, tblPr)
remove_borders(tblPr, tag="w:tblBorders")

left_cell  = table.cell(0, 0)
right_cell = table.cell(0, 1)

# Largeurs : ~8 cm (émetteur) | ~8 cm (destinataire)
set_cell_width(left_cell,  8.0)
set_cell_width(right_cell, 8.0)

# Supprimer les bordures sur chaque cellule
for cell in [left_cell, right_cell]:
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    remove_borders(tcPr, tag="w:tcBorders")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

# Vider les paragraphes vides créés par défaut
for cell in [left_cell, right_cell]:
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)

# Colonne gauche – expéditrice
for text, bold, size, sa in [
    ("Louiza HADID",          True,  12, 4),
    ("3 rue de Ventspils",    False, 11, 2),
    ("56100 Lorient",         False, 11, 2),
    ("(+33) 6 50 37 56 47",  False, 11, 2),
    ("louizahdd@gmail.com",   False, 11, 2),
]:
    add_cell_para(left_cell, text, bold=bold, size=size, sa=sa)

# Colonne droite – destinataire (aligné à gauche dans sa cellule)
for text, bold, size, sa in [
    ("ISFEC Bretagne",                  True,  11, 3),
    ("Site d'Arradon",                  False, 11, 2),
    ("3 allée des Fougères, BP 25",     False, 11, 2),
    ("56610 Arradon",                   False, 11, 2),
]:
    add_cell_para(right_cell, text, bold=bold, size=size, sa=sa)

# ── Espace + Date ─────────────────────────────────────────────────────────
para(doc, sb=10, sa=0)
para(doc, "Lorient, le 1er mars 2026",
     align=WD_ALIGN_PARAGRAPH.RIGHT, sb=2, sa=14)

# ── Objet ─────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after  = Pt(14)
r1 = p.add_run("Objet : ")
font(r1, bold=True)
r2 = p.add_run("Candidature au Master MEEF – Mention 1er degré")
font(r2, bold=True)

# ── Salutation ────────────────────────────────────────────────────────────
para(doc, "Madame, Monsieur,", sb=0, sa=12)

# ════════════════════════════════════════════════════════════════════════════
#  CORPS
# ════════════════════════════════════════════════════════════════════════════

corps = [

    # § 1 ─ Présentation et choix de l'ISFEC d'Arradon
    (
        "C'est avec un vif intérêt que je vous adresse ma candidature pour le Master MEEF, "
        "mention 1er degré, proposé par l'ISFEC Bretagne sur son site d'Arradon. Affilié à "
        "l'Université Catholique de l'Ouest pour la délivrance du diplôme, cet établissement "
        "s'inscrit pleinement dans le réseau national des ISFEC et se distingue par un "
        "accompagnement de proximité au sein de promotions à effectifs réduits. C'est ce cadre "
        "de formation, à la fois rigoureux et bienveillant, ancré dans les valeurs de "
        "l'enseignement catholique, qui en fait mon premier choix dans la perspective de "
        "préparer le CRPE et d'enseigner dans le premier degré."
    ),

    # § 2 ─ Parcours académique
    (
        "Mon parcours en Sciences de l'éducation à l'Université Rennes 2 m'a permis de "
        "construire des bases solides en pédagogie, en psychologie du développement et en "
        "didactique. Ces apprentissages ont progressivement conforté ma conviction : c'est dans "
        "l'enseignement du premier degré que je veux m'investir sur le long terme. La "
        "préparation au CRPE intégrée à votre Master constitue pour moi une étape décisive "
        "pour accéder à ce métier dans les meilleures conditions."
    ),

    # § 3 ─ Expériences professionnelles (sans tiret cadratin)
    (
        "Sur le plan pratique, mes expériences au contact des enfants ont renforcé cette "
        "orientation. En tant qu'animatrice périscolaire dans le quartier prioritaire de "
        "Villejean à Rennes, j'ai appris à adapter mon approche à des publics diversifiés et à "
        "faire preuve de patience et de bienveillance dans des situations parfois complexes. "
        "J'ai également exercé comme tutrice de français auprès d'apprenants étrangers, ce qui "
        "m'a amenée à travailler la reformulation, la différenciation pédagogique et l'écoute "
        "active, des compétences directement transposables dans l'exercice du métier "
        "d'enseignante."
    ),

    # § 4 ─ Conclusion et motivation
    (
        "Ces expériences m'ont confirmé que l'enseignement va bien au-delà de la transmission "
        "de savoirs : il s'agit avant tout d'accompagner chaque élève dans son développement, "
        "de l'encourager et d'adapter sa pratique à ses besoins. C'est dans cette perspective "
        "que je souhaite m'investir pleinement dans votre formation, convaincue que l'ISFEC "
        "d'Arradon m'offrira les outils théoriques et pratiques pour exercer ce métier avec "
        "rigueur et engagement."
    ),
]

for texte in corps:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(10)
    p.paragraph_format.first_line_indent = Cm(1)
    r = p.add_run(texte)
    font(r)

# ── Formule de politesse ──────────────────────────────────────────────────
for texte, sa in [
    (
        "Dans l'attente d'une réponse favorable, je me tiens à votre disposition "
        "pour tout entretien ou renseignement complémentaire.",
        10,
    ),
    (
        "Veuillez recevoir, Madame, Monsieur, l'expression de mes salutations distinguées.",
        30,
    ),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(sa)
    r = p.add_run(texte)
    font(r)

# ── Signature ─────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after  = Pt(0)
r = p.add_run("Louiza Hadid")
font(r, bold=True)

# ── Sauvegarde ────────────────────────────────────────────────────────────
output = "/home/user/Lettre-de-motivation/Lettre_motivation_ISFEC_Arradon_M2E.docx"
doc.save(output)
print(f"Fichier généré : {output}")
